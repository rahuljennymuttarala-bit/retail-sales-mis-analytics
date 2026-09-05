import pandas as pd

file_path = "C:/Users/rahul/OneDrive/Desktop/portfolio/project 1_Retail _sales/executive_sales_retail_mis_raw_dataset.xlsx"

excel = pd.ExcelFile(file_path)

print("Available Sheets:")
print(excel.sheet_names)

# ==============================
# DATA PROFILING
# ==============================

for sheet in excel.sheet_names:

    df = pd.read_excel(file_path, sheet_name=sheet)

    print("\n" + "=" * 60)
    print("SHEET:", sheet)
    print("=" * 60)

    print("Rows:", df.shape[0])
    print("Columns:", df.shape[1])

    print("\nColumn Names:")
    print(df.columns.tolist())

    print("\nMissing Values:")
    print(df.isnull().sum())

    print("\nDuplicate Rows:")
    print(df.duplicated().sum())


# ==============================
# DUPLICATE INVESTIGATION
# ==============================

customers = pd.read_excel(file_path, sheet_name="Customers")

print("\nCustomer Duplicate Record:")
print(customers[customers.duplicated(keep=False)])


products = pd.read_excel(file_path, sheet_name="Products")

print("\nProduct Duplicate Record:")
print(products[products.duplicated(keep=False)])


stores = pd.read_excel(file_path, sheet_name="Stores")

print("\nStore Duplicate Record:")
print(stores[stores.duplicated(keep=False)])


orders = pd.read_excel(file_path, sheet_name="Orders")

print("\nOrder Duplicate Record:")
print(orders[orders.duplicated(keep=False)])

# ==============================
# REMOVE TRUE DUPLICATES
# ==============================

customers_clean = customers.drop_duplicates()
products_clean = products.drop_duplicates()
stores_clean = stores.drop_duplicates()
orders_clean = orders.drop_duplicates()

print("\nDuplicate count after cleaning:")

print("Customers:", customers_clean.duplicated().sum())
print("Products:", products_clean.duplicated().sum())
print("Stores:", stores_clean.duplicated().sum())
print("Orders:", orders_clean.duplicated().sum())

# ==============================
# MISSING VALUE INVESTIGATION
# ==============================

print("\nCustomers - Missing Email Records:")
print(customers[customers["Email"].isna()])

print("\nCustomers - Missing Phone Records:")
print(customers[customers["Phone"].isna()])

order_details = pd.read_excel(
    file_path,
    sheet_name="Order_Details"
)

print("\nMissing Fulfillment Date:")
print(
    order_details[
        order_details["Fulfillment_Date"].isna()
    ][
        ["Order_ID", "Product_ID", "Fulfillment_Status", "Fulfillment_Date"]
    ].head(20)
)

print("\nMissing Promo Code:")
print(
    order_details[
        order_details["Promo_Code"].isna()
    ][
        ["Order_ID", "Product_ID", "Discount_Pct", "Promo_Code"]
    ].head(20)
)

# ==============================
# HANDLE MISSING VALUES
# ==============================

# Keep missing Email and Phone as they are.
# We should not invent customer contact information.

# Fulfillment_Date:
# Keep missing dates because cancelled orders were not fulfilled.

# Promo_Code:
# Replace missing promo codes with "No Promo".

order_details_clean = order_details.copy()

order_details_clean["Promo_Code"] = (
    order_details_clean["Promo_Code"]
    .fillna("No Promo")
)

print("\nMissing values after handling Promo_Code:")
print(order_details_clean["Promo_Code"].isnull().sum())

# ==============================
# CHECK CATEGORICAL VALUES
# ==============================

categorical_columns = {
    "Customers": ["Gender", "Loyalty_Tier"],
    "Retail_Transactions_10K": ["Region", "Sales_Channel", "Category", "Return_Flag"],
    "Products": ["Category", "Subcategory", "Brand", "Product_Status"],
    "Stores": ["Region", "Store_Format", "Store_Status"],
    "Orders": ["Sales_Channel", "Payment_Method", "Order_Status"],
    "Order_Details": ["Fulfillment_Status", "Promo_Code"],
    "Returns": ["Return_Reason", "Return_Status"]
}

for sheet, columns in categorical_columns.items():

    df = pd.read_excel(file_path, sheet_name=sheet)

    print("\n" + "=" * 60)
    print("SHEET:", sheet)
    print("=" * 60)

    for column in columns:
        print("\n", column)
        print(df[column].value_counts(dropna=False))


        # ==============================
# STANDARDIZE TEXT VALUES
# ==============================

customers_clean["Loyalty_Tier"] = (
    customers_clean["Loyalty_Tier"]
    .astype(str)
    .str.strip()
    .str.title()
)

customers_clean["Gender"] = (
    customers_clean["Gender"]
    .astype(str)
    .str.strip()
    .str.title()
)

order_details_clean["Fulfillment_Status"] = (
    order_details_clean["Fulfillment_Status"]
    .astype(str)
    .str.strip()
    .str.title()
)

products_clean["Product_Status"] = (
    products_clean["Product_Status"]
    .astype(str)
    .str.strip()
    .str.title()
)

stores_clean["Store_Status"] = (
    stores_clean["Store_Status"]
    .astype(str)
    .str.strip()
    .str.title()
)

orders_clean["Order_Status"] = (
    orders_clean["Order_Status"]
    .astype(str)
    .str.strip()
    .str.title()
)

print("\nCustomer Loyalty Tiers:")
print(customers_clean["Loyalty_Tier"].value_counts())

print("\nCustomer Gender:")
print(customers_clean["Gender"].value_counts())

print("\nOrder Status:")
print(orders_clean["Order_Status"].value_counts())

print("\nProduct Status:")
print(products_clean["Product_Status"].value_counts())

print("\nStore Status:")
print(stores_clean["Store_Status"].value_counts())

print("\nFulfillment Status:")
print(order_details_clean["Fulfillment_Status"].value_counts())


# ==============================
# DATA TYPE VALIDATION
# ==============================

# Check data types for each cleaned dataset

print("\nCustomers Data Types:")
print(customers_clean.dtypes)

print("\nProducts Data Types:")
print(products_clean.dtypes)

print("\nStores Data Types:")
print(stores_clean.dtypes)

print("\nOrders Data Types:")
print(orders_clean.dtypes)

print("\nOrder Details Data Types:")
print(order_details_clean.dtypes)

print("\nDate Column Types:")

print("Customer Join Date:",
      customers_clean["Join_Date"].dtype)

print("Product Launch Date:",
      products_clean["Launch_Date"].dtype)

print("Store Opening Date:",
      stores_clean["Opening_Date"].dtype)

print("Order Date:",
      orders_clean["Order_Date"].dtype)

print("Fulfillment Date:",
      order_details_clean["Fulfillment_Date"].dtype)

# ==============================
# BUSINESS RULE VALIDATION
# ==============================

print("\n========== BUSINESS VALIDATION ==========")

# 1. Quantity should be greater than 0
print("\nInvalid Quantity:")
print(
    order_details_clean[
        order_details_clean["Quantity"] <= 0
    ][["Order_Detail_ID", "Product_ID", "Quantity"]]
)

# 2. Unit price should not be negative
print("\nInvalid Unit Price:")
print(
    order_details_clean[
        order_details_clean["Unit_Price"] < 0
    ][["Order_Detail_ID", "Product_ID", "Unit_Price"]]
)

# 3. Discount percentage should be between 0 and 100
print("\nInvalid Discount:")
print(
    order_details_clean[
        (order_details_clean["Discount_Pct"] < 0) |
        (order_details_clean["Discount_Pct"] > 100)
    ][["Order_Detail_ID", "Discount_Pct"]]
)

# 4. Tax percentage should be between 0 and 100
print("\nInvalid Tax:")
print(
    order_details_clean[
        (order_details_clean["Tax_Pct"] < 0) |
        (order_details_clean["Tax_Pct"] > 100)
    ][["Order_Detail_ID", "Tax_Pct"]]
)

# 5. Sales values should not be negative
print("\nNegative Net Sales:")
print(
    order_details_clean[
        order_details_clean["Net_Sales"] < 0
    ][["Order_Detail_ID", "Net_Sales"]]
)

# 6. Check calculation: Gross Sales - Discount = Net Sales
order_details_clean["Calculated_Net_Sales"] = (
    order_details_clean["Gross_Sales"]
    - order_details_clean["Discount_Amount"]
)

difference = (
    order_details_clean["Calculated_Net_Sales"]
    - order_details_clean["Net_Sales"]
).abs()

print("\nIncorrect Net Sales Calculations:")
print((difference > 0.01).sum())

# ==============================
# CREATE FINAL CLEAN DATASETS
# ==============================

retail_clean = pd.read_excel(
    file_path,
    sheet_name="Retail_Transactions_10K"
)

# Remove duplicate transactions if any
retail_clean = retail_clean.drop_duplicates()

# Clean customer data
customers_clean = customers.drop_duplicates()

# Clean product data
products_clean = products.drop_duplicates()

# Clean store data
stores_clean = stores.drop_duplicates()

# Clean orders
orders_clean = orders.drop_duplicates()

# Clean order details
order_details_clean = order_details.drop_duplicates()

# Handle missing promo codes
order_details_clean["Promo_Code"] = (
    order_details_clean["Promo_Code"]
    .fillna("No Promo")
)

# Standardize text
customers_clean["Loyalty_Tier"] = (
    customers_clean["Loyalty_Tier"]
    .astype(str)
    .str.strip()
    .str.title()
)

customers_clean["Gender"] = (
    customers_clean["Gender"]
    .astype(str)
    .str.strip()
    .str.title()
)

products_clean["Product_Status"] = (
    products_clean["Product_Status"]
    .astype(str)
    .str.strip()
    .str.title()
)

stores_clean["Store_Status"] = (
    stores_clean["Store_Status"]
    .astype(str)
    .str.strip()
    .str.title()
)

orders_clean["Order_Status"] = (
    orders_clean["Order_Status"]
    .astype(str)
    .str.strip()
    .str.title()
)

order_details_clean["Fulfillment_Status"] = (
    order_details_clean["Fulfillment_Status"]
    .astype(str)
    .str.strip()
    .str.title()
)

print("\nFinal Clean Dataset Sizes:")
print("Retail Transactions:", retail_clean.shape)
print("Customers:", customers_clean.shape)
print("Products:", products_clean.shape)
print("Stores:", stores_clean.shape)
print("Orders:", orders_clean.shape)
print("Order Details:", order_details_clean.shape)

# ==============================
# SAVE CLEANED DATA
# ==============================

output_file = "C:/Users/rahul/OneDrive/Desktop/portfolio/project 1_Retail _sales/cleaned_data/retail_cleaned_data.xlsx"

with pd.ExcelWriter(output_file, engine="openpyxl") as writer:

    retail_clean.to_excel(
        writer,
        sheet_name="Retail_Transactions",
        index=False
    )

    customers_clean.to_excel(
        writer,
        sheet_name="Customers",
        index=False
    )

    products_clean.to_excel(
        writer,
        sheet_name="Products",
        index=False
    )

    stores_clean.to_excel(
        writer,
        sheet_name="Stores",
        index=False
    )

    orders_clean.to_excel(
        writer,
        sheet_name="Orders",
        index=False
    )

    order_details_clean.to_excel(
        writer,
        sheet_name="Order_Details",
        index=False
    )

print("\nCleaned Excel file created successfully!")


# ==============================
# STEP 11 - SALES EDA
# ==============================

print("\n========== SALES OVERVIEW ==========")

# Total Sales
total_sales = retail_clean["Net_Sales"].sum()

# Total Transactions
total_transactions = retail_clean["Transaction_ID"].nunique()

# Total Orders
total_orders = retail_clean["Order_ID"].nunique()

# Total Quantity Sold
total_quantity = retail_clean["Quantity"].sum()

# Average Transaction Value
average_transaction_value = (
    total_sales / total_transactions
)

print("Total Sales:", total_sales)
print("Total Transactions:", total_transactions)
print("Total Orders:", total_orders)
print("Total Quantity Sold:", total_quantity)
print("Average Transaction Value:", average_transaction_value)

# ==============================
# REGION PERFORMANCE
# ==============================

region_sales = (
    retail_clean
    .groupby("Region")
    .agg(
        Total_Sales=("Net_Sales", "sum"),
        Transactions=("Transaction_ID", "nunique"),
        Quantity=("Quantity", "sum")
    )
    .sort_values("Total_Sales", ascending=False)
)

print("\n========== REGION PERFORMANCE ==========")
print(region_sales)


# ==============================
# CATEGORY PERFORMANCE
# ==============================

category_sales = (
    retail_clean
    .groupby("Category")
    .agg(
        Total_Sales=("Net_Sales", "sum"),
        Transactions=("Transaction_ID", "nunique"),
        Quantity=("Quantity", "sum")
    )
    .sort_values("Total_Sales", ascending=False)
)

print("\n========== CATEGORY PERFORMANCE ==========")
print(category_sales)

# ==============================
# MONTHLY SALES TREND
# ==============================

retail_clean["Order_Date"] = pd.to_datetime(
    retail_clean["Order_Date"],
    format="mixed",
    dayfirst=True
)
retail_clean["Month"] = (
    retail_clean["Order_Date"]
    .dt.to_period("M")
    .astype(str)
)

monthly_sales = (
    retail_clean
    .groupby("Month")
    .agg(
        Total_Sales=("Net_Sales", "sum"),
        Transactions=("Transaction_ID", "nunique"),
        Quantity=("Quantity", "sum")
    )
    .reset_index()
)

print("\n========== MONTHLY SALES ==========")
print(monthly_sales)

monthly_sales["MoM_Growth_%"] = (
    monthly_sales["Total_Sales"]
    .pct_change()
    * 100
)

print("\n========== MONTHLY SALES WITH GROWTH ==========")
print(monthly_sales)

# ==============================
# STORE PERFORMANCE & TARGETS
# ==============================

sales_targets = pd.read_excel(
    file_path,
    sheet_name="Sales_Targets"
)

print("\nSales Targets:")
print(sales_targets.head())

sales_targets["Month"] = pd.to_datetime(
    sales_targets["Month"],
    format="mixed"
).dt.to_period("M").astype(str)

store_monthly_sales = (
    retail_clean
    .groupby(["Month", "Store_ID"])
    .agg(
        Actual_Sales=("Net_Sales", "sum"),
        Actual_Orders=("Order_ID", "nunique")
    )
    .reset_index()
)

store_performance = pd.merge(
    store_monthly_sales,
    sales_targets,
    on=["Month", "Store_ID"],
    how="left"
)

store_performance["Sales_Achievement_%"] = (
    store_performance["Actual_Sales"]
    / store_performance["Target_Sales"]
    * 100
)

store_performance["Sales_Variance"] = (
    store_performance["Actual_Sales"]
    - store_performance["Target_Sales"]
)

store_performance["Order_Achievement_%"] = (
    store_performance["Actual_Orders"]
    / store_performance["Target_Orders"]
    * 100
)

print("\n========== STORE PERFORMANCE ==========")

print(
    store_performance[
        [
            "Month",
            "Store_ID",
            "Actual_Sales",
            "Target_Sales",
            "Sales_Achievement_%",
            "Sales_Variance",
            "Actual_Orders",
            "Target_Orders",
            "Order_Achievement_%"
        ]
    ].head(20)
)

underperforming = store_performance[
    store_performance["Sales_Achievement_%"] < 100
]

print("\n========== UNDERPERFORMING STORE-MONTHS ==========")
print(
    underperforming[
        [
            "Month",
            "Store_ID",
            "Actual_Sales",
            "Target_Sales",
            "Sales_Achievement_%",
            "Sales_Variance"
        ]
    ]
    .sort_values("Sales_Achievement_%")
    .head(20)
)

top_performers = store_performance[
    store_performance["Sales_Achievement_%"] >= 100
]

print("\n========== TARGET ACHIEVERS ==========")
print(
    top_performers[
        [
            "Month",
            "Store_ID",
            "Actual_Sales",
            "Target_Sales",
            "Sales_Achievement_%",
            "Sales_Variance"
        ]
    ]
    .sort_values("Sales_Achievement_%", ascending=False)
    .head(20)
)

top_performers = store_performance[
    store_performance["Sales_Achievement_%"] >= 100
]

print("\n========== TARGET ACHIEVERS ==========")
print(
    top_performers[
        [
            "Month",
            "Store_ID",
            "Actual_Sales",
            "Target_Sales",
            "Sales_Achievement_%",
            "Sales_Variance"
        ]
    ]
    .sort_values("Sales_Achievement_%", ascending=False)
    .head(20)
)

# ==============================
# STEP 14 - MONTHLY CHANNEL ANALYSIS
# ==============================

monthly_channel = (
    retail_clean
    .groupby(["Month", "Sales_Channel"])
    .agg(
        Sales=("Net_Sales", "sum"),
        Transactions=("Transaction_ID", "nunique")
    )
    .reset_index()
)

print("\n========== MONTHLY CHANNEL PERFORMANCE ==========")
print(monthly_channel)

# ==============================
# MONTHLY REGION ANALYSIS
# ==============================

monthly_region = (
    retail_clean
    .groupby(["Month", "Region"])
    .agg(
        Sales=("Net_Sales", "sum"),
        Transactions=("Transaction_ID", "nunique")
    )
    .reset_index()
)

print("\n========== MONTHLY REGION PERFORMANCE ==========")
print(monthly_region)

# ==============================
# STEP 15 - CUSTOMER ANALYSIS
# ==============================

customer_sales = (
    retail_clean
    .groupby("Customer_ID")
    .agg(
        Total_Sales=("Net_Sales", "sum"),
        Transactions=("Transaction_ID", "nunique"),
        Quantity=("Quantity", "sum")
    )
    .reset_index()
    .sort_values("Total_Sales", ascending=False)
)

print("\n========== TOP 10 CUSTOMERS ==========")
print(customer_sales.head(10))

# ==============================
# LOYALTY TIER PERFORMANCE
# ==============================

customer_loyalty = customers_clean[
    ["Customer_ID", "Loyalty_Tier"]
]

loyalty_analysis = pd.merge(
    customer_sales,
    customer_loyalty,
    on="Customer_ID",
    how="left"
)

loyalty_performance = (
    loyalty_analysis
    .groupby("Loyalty_Tier")
    .agg(
        Total_Sales=("Total_Sales", "sum"),
        Customers=("Customer_ID", "nunique"),
        Transactions=("Transactions", "sum")
    )
    .reset_index()
    .sort_values("Total_Sales", ascending=False)
)

print("\n========== LOYALTY TIER PERFORMANCE ==========")
print(loyalty_performance)

# ==============================
# SALES CHANNEL PERFORMANCE
# ==============================

channel_performance = (
    retail_clean
    .groupby("Sales_Channel")
    .agg(
        Total_Sales=("Net_Sales", "sum"),
        Transactions=("Transaction_ID", "nunique"),
        Quantity_Sold=("Quantity", "sum")
    )
    .reset_index()
    .sort_values("Total_Sales", ascending=False)
)

channel_performance["Sales_Contribution_%"] = (
    channel_performance["Total_Sales"]
    / channel_performance["Total_Sales"].sum()
    * 100
)

print("\n========== SALES CHANNEL PERFORMANCE ==========")
print(channel_performance)

# ==============================
# STEP 17 - RETURNS ANALYSIS
# ==============================

returns = pd.read_excel(
    file_path,
    sheet_name="Returns"
)

total_returns = returns["Return_ID"].nunique()

return_rate = (
    total_returns
    / total_transactions
    * 100
)

total_refund = returns["Refund_Amount"].sum()

print("\n========== RETURNS OVERVIEW ==========")
print("Total Returns:", total_returns)
print("Return Rate %:", return_rate)
print("Total Refund Amount:", total_refund)

return_reasons = (
    returns
    .groupby("Return_Reason")
    .agg(
        Return_Count=("Return_ID", "nunique"),
        Refund_Amount=("Refund_Amount", "sum")
    )
    .reset_index()
    .sort_values("Return_Count", ascending=False)
)

print("\n========== RETURN REASONS ==========")
print(return_reasons)




# ============================================
# STEP 18 - EXPORT MIS ANALYSIS TO EXCEL
# ============================================

import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# --------------------------------------------
# OUTPUT FOLDER
# --------------------------------------------

output_folder = (
    "C:/Users/rahul/OneDrive/Desktop/portfolio/"
    "project 1_Retail _sales/excel"
)

os.makedirs(output_folder, exist_ok=True)

mis_file = os.path.join(
    output_folder,
    "Retail_MIS_Analysis.xlsx"
)

# --------------------------------------------
# CREATE PRODUCT PERFORMANCE
# --------------------------------------------

product_performance = (
    retail_clean
    .groupby("Product_ID")
    .agg(
        Total_Sales=("Net_Sales", "sum"),
        Quantity_Sold=("Quantity", "sum"),
        Transactions=("Transaction_ID", "nunique")
    )
    .reset_index()
    .sort_values("Total_Sales", ascending=False)
)

# --------------------------------------------
# CREATE CATEGORY PERFORMANCE
# --------------------------------------------

category_performance = (
    retail_clean
    .groupby("Category")
    .agg(
        Total_Sales=("Net_Sales", "sum"),
        Quantity_Sold=("Quantity", "sum"),
        Transactions=("Transaction_ID", "nunique")
    )
    .reset_index()
    .sort_values("Total_Sales", ascending=False)
)

category_performance["Sales_Contribution_%"] = (
    category_performance["Total_Sales"]
    / category_performance["Total_Sales"].sum()
    * 100
)

# --------------------------------------------
# RETURNS
# --------------------------------------------

total_returns = returns["Return_ID"].nunique()

total_refund = returns["Refund_Amount"].sum()

return_rate = (
    total_returns
    / total_transactions
    * 100
)

# --------------------------------------------
# TARGET ACHIEVEMENT
# --------------------------------------------

overall_target_sales = (
    sales_targets["Target_Sales"].sum()
)

overall_actual_sales = (
    store_performance["Actual_Sales"].sum()
)

overall_target_achievement = (
    overall_actual_sales
    / overall_target_sales
    * 100
)

# --------------------------------------------
# EXECUTIVE KPI TABLE
# --------------------------------------------

executive_kpis = pd.DataFrame({

    "KPI": [
        "Total Sales",
        "Total Orders",
        "Total Transactions",
        "Total Quantity Sold",
        "Average Transaction Value",
        "Average Order Value",
        "Total Returns",
        "Return Rate %",
        "Total Refund Amount",
        "Sales Target Achievement %"
    ],

    "Value": [
        total_sales,
        total_orders,
        total_transactions,
        total_quantity,
        average_transaction_value,
        total_sales / total_orders,
        total_returns,
        return_rate,
        total_refund,
        overall_target_achievement
    ]
})

# --------------------------------------------
# EXPORT TO EXCEL
# --------------------------------------------

with pd.ExcelWriter(
    mis_file,
    engine="openpyxl"
) as writer:

    executive_kpis.to_excel(
        writer,
        sheet_name="Executive KPIs",
        index=False
    )

    monthly_sales.to_excel(
        writer,
        sheet_name="Monthly Sales",
        index=False
    )

    region_sales.reset_index().to_excel(
        writer,
        sheet_name="Region Performance",
        index=False
    )

    category_performance.to_excel(
        writer,
        sheet_name="Category Performance",
        index=False
    )

    store_performance.to_excel(
        writer,
        sheet_name="Store Performance",
        index=False
    )

    product_performance.to_excel(
        writer,
        sheet_name="Product Performance",
        index=False
    )

    customer_sales.to_excel(
        writer,
        sheet_name="Customer Analysis",
        index=False
    )

    loyalty_performance.to_excel(
        writer,
        sheet_name="Loyalty Analysis",
        index=False
    )

    channel_performance.to_excel(
        writer,
        sheet_name="Channel Analysis",
        index=False
    )

    return_reasons.to_excel(
        writer,
        sheet_name="Returns Analysis",
        index=False
    )

# --------------------------------------------
# FORMAT WORKBOOK
# --------------------------------------------

workbook = load_workbook(mis_file)

for worksheet in workbook.worksheets:

    # Header formatting
    for cell in worksheet[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Freeze first row
    worksheet.freeze_panes = "A2"

    # Auto-adjust column width
    for column in worksheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35
        )

# --------------------------------------------
# SAVE
# --------------------------------------------

workbook.save(mis_file)

print("\n============================================")
print("MIS EXCEL REPORT CREATED SUCCESSFULLY!")
print("============================================")
print("File:", mis_file)